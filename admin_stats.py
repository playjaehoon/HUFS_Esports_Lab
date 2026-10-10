"""Staff-only aggregate dashboard and spreadsheet export."""
from datetime import date, timedelta
from io import BytesIO

from flask import render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import distinct

from auth_helpers import admin_required
from booking import RuleError, korea_now
from models import Reservation, Student, db


STATUS_LABELS = {'active': '예약 완료/이용 중', 'completed': '이용 완료',
                 'cancelled': '취소', 'no_show': '미방문(No-show)'}
STATUS_COLORS = ('#2474aa', '#23a6b8', '#91aabb', '#c96c71')


def selected_period():
    scope = request.args.get('scope', 'range')
    if scope == 'all':
        return scope, None, None
    if scope != 'range':
        raise RuleError('통계 조회 범위를 확인해 주세요.')
    today = korea_now().date()
    try:
        start = date.fromisoformat(request.args.get('start', (today - timedelta(days=29)).isoformat()))
        end = date.fromisoformat(request.args.get('end', today.isoformat()))
    except (ValueError, TypeError):
        raise RuleError('조회 날짜를 확인해 주세요.') from None
    if start > end or (end - start).days > 365:
        raise RuleError('조회 기간은 최대 366일이며 시작일이 종료일보다 늦을 수 없습니다.')
    return scope, start, end


def safe_cell(value):
    """Prevent Excel from executing student-entered strings as formulas."""
    if isinstance(value, str) and value and value[0] in '=+-@\t\r\n':
        return "'" + value
    return value


def aggregates(scope, start, end):
    conditions = [Reservation.date.between(start.isoformat(), end.isoformat())] if scope == 'range' else []
    base = db.select(Reservation).where(*conditions)
    status_counts = dict(db.session.execute(
        db.select(Reservation.status, db.func.count()).where(*conditions)
        .group_by(Reservation.status)).all())
    daily_counts = dict(db.session.execute(
        db.select(Reservation.date, db.func.count()).where(*conditions)
        .group_by(Reservation.date)).all())
    if scope == 'range':
        day_count = (end - start).days + 1
        daily = [((start + timedelta(days=n)).isoformat(),
                  daily_counts.get((start + timedelta(days=n)).isoformat(), 0)) for n in range(day_count)]
    else:
        observed = sorted(daily_counts)
        first_day = date.fromisoformat(observed[0]) if observed else None
        last_day = date.fromisoformat(observed[-1]) if observed else None
        day_count = (last_day - first_day).days + 1 if observed else 0
        daily = [((first_day + timedelta(days=n)).isoformat(),
                  daily_counts.get((first_day + timedelta(days=n)).isoformat(), 0)) for n in range(day_count)]
    if scope == 'all' or day_count > 90:
        monthly = {}
        for day, count in daily_counts.items():
            month = day[:7]
            monthly[month] = monthly.get(month, 0) + count
        first_month = start if scope == 'range' else (date.fromisoformat(daily[0][0]) if daily else None)
        last_month = end if scope == 'range' else (date.fromisoformat(daily[-1][0]) if daily else None)
        trend = []
        if first_month and last_month:
            year, month = first_month.year, first_month.month
            while (year, month) <= (last_month.year, last_month.month):
                label = f'{year:04d}-{month:02d}'
                trend.append((label, monthly.get(label, 0)))
                month += 1
                if month == 13:
                    year, month = year + 1, 1
        trend_unit = '월별'
    else:
        trend = daily
        trend_unit = '일별'
    peak = max((count for _, count in trend), default=0) or 1
    points = [(round(24 + i * 592 / max(len(trend) - 1, 1)), round(154 - count * 120 / peak))
              for i, (_, count) in enumerate(trend)]
    point_text = ' '.join(f'{x},{y}' for x, y in points)
    area_text = f'24,154 {point_text} {points[-1][0]},154' if points else ''
    department_counts = db.session.execute(
        db.select(Student.department, db.func.count()).group_by(Student.department)
        .order_by(db.func.count().desc(), Student.department)).all()
    booked_minutes = db.session.scalar(
        db.select(db.func.sum(Reservation.end_minute - Reservation.start_minute))
        .where(*conditions, Reservation.status != 'cancelled')) or 0
    attended = db.session.scalar(db.select(db.func.count()).select_from(Reservation)
                                 .where(*conditions, Reservation.is_attended.is_(True))) or 0
    visitors = db.session.scalar(db.select(db.func.count(distinct(Reservation.student_pk)))
                                 .where(*conditions, Reservation.is_attended.is_(True))) or 0
    total_bookings = sum(status_counts.values())
    valid_bookings = total_bookings - status_counts.get('cancelled', 0)
    boundaries = []
    running = 0
    for key in STATUS_LABELS:
        running += status_counts.get(key, 0) / (total_bookings or 1) * 100
        boundaries.append(round(running, 2))
    donut_background = ('conic-gradient(' + ', '.join(
        f'{color} {0 if i == 0 else boundaries[i-1]}% {boundaries[i]}%'
        for i, color in enumerate(STATUS_COLORS)) + ')') if total_bookings else '#e2ebf0'
    return {
        'total_students': db.session.scalar(db.select(db.func.count()).select_from(Student)) or 0,
        'active_students': db.session.scalar(db.select(db.func.count()).select_from(Student)
                                             .where(Student.archived.is_(False))) or 0,
        'total_bookings': total_bookings,
        'status_counts': [(STATUS_LABELS[key], status_counts.get(key, 0)) for key in STATUS_LABELS],
        'status_colors': STATUS_COLORS, 'donut_background': donut_background,
        'daily': daily, 'trend': trend, 'trend_unit': trend_unit,
        'trend_points': point_text, 'trend_area': area_text, 'trend_peak': peak,
        'trend_first_point': points[0] if points else None,
        'departments': [(name or '미입력', count) for name, count in department_counts],
        'booked_hours': round(booked_minutes / 60, 1), 'attended': attended,
        'visitors': visitors, 'attendance_rate': round(attended / valid_bookings * 100, 1) if valid_bookings else 0,
        'cancelled': status_counts.get('cancelled', 0),
        'reservations': base,
    }


def sheet(workbook, name, headings, rows):
    tab = workbook.create_sheet(name)
    tab.append(headings)
    for cell in tab[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='002C5F')
    for row in rows:
        tab.append([safe_cell(item) for item in row])
    tab.freeze_panes = 'A2'
    tab.auto_filter.ref = tab.dimensions
    return tab


def register_stats_routes(app):
    @app.route('/admin/stats')
    @admin_required
    def admin_stats():
        scope, start, end = selected_period()
        stats = aggregates(scope, start, end)
        return render_template('admin_stats.html', stats=stats, scope=scope, start=start, end=end)

    @app.route('/admin/stats/export.xlsx')
    @admin_required
    def admin_stats_export():
        scope, start, end = selected_period()
        stats = aggregates(scope, start, end)
        workbook = Workbook()
        workbook.remove(workbook.active)
        sheet(workbook, '요약', ['항목', '값'], [
            ('조회 범위', '전체 기간' if scope == 'all' else '기간 설정'),
            ('조회 시작', start.isoformat() if start else '전체'),
            ('조회 종료', end.isoformat() if end else '전체'),
            ('회원가입 학생 수 (누적)', stats['total_students']), ('활성 회원', stats['active_students']),
            ('예약 수 (취소 포함)', stats['total_bookings']), ('방문 학생 수 (중복 제거)', stats['visitors']),
            ('방문 횟수', stats['attended']), ('예약 시간 (취소 제외)', stats['booked_hours']),
            ('방문율 (%)', stats['attendance_rate']), *stats['status_counts']])
        sheet(workbook, '일별 예약', ['날짜', '예약 건수'], stats['daily'])
        sheet(workbook, '학과별 회원', ['학과', '회원 수'], stats['departments'])
        students = db.session.scalars(db.select(Student).order_by(Student.student_number))
        sheet(workbook, '학생', ['학번', '이름', '학과', '계정 상태'],
              ((item.student_number, item.name, item.department or '',
                '비활성화' if item.archived else '활성') for item in students))
        reservations = db.session.scalars(stats['reservations'].order_by(Reservation.date, Reservation.id))
        sheet(workbook, '예약', ['예약 번호', '날짜', '학번', '이름', '좌석', '시작', '종료', '상태', '방문 확인'],
              ((item.id, item.date, item.student_id, item.student_name, item.seat_number,
                f'{item.start_minute // 60:02d}:{item.start_minute % 60:02d}',
                f'{item.end_minute // 60:02d}:{item.end_minute % 60:02d}',
                STATUS_LABELS.get(item.status, item.status), '예' if item.is_attended else '아니요')
               for item in reservations))
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        period_name = 'all' if scope == 'all' else f'{start.isoformat()}-{end.isoformat()}'
        return send_file(output, as_attachment=True,
                         download_name=f'evnia-statistics-{period_name}.xlsx',
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
