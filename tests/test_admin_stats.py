from io import BytesIO

from openpyxl import load_workbook

from conftest import login, reserve
from models import Reservation, Student, db


def test_student_sorting_and_stats_export_are_admin_only(app):
    admin = login(app, admin=True)
    student = login(app)
    assert student.get('/admin/stats').status_code == 403
    assert student.get('/admin/stats/export.xlsx').status_code == 403
    with app.app_context():
        person = db.session.get(Student, 1)
        person.name = '=HYPERLINK("https://example.invalid")'
        person.department = '글로벌스포츠산업학부'
        db.session.commit()
    reserve(student)
    with app.app_context():
        first = db.session.scalar(db.select(Reservation))
        first.is_attended = True
        db.session.add(Reservation(student_pk=1, student_id='202600001', student_name=first.student_name,
                                   date='2026-10-01', start_minute=600, end_minute=660,
                                   seat_number=1, status='completed', is_attended=True))
        db.session.commit()
    listing = admin.get('/admin/students?sort=department&direction=desc')
    assert listing.status_code == 200
    assert '글로벌스포츠산업학부' in listing.get_data(as_text=True)
    assert 'sort=department' in listing.get_data(as_text=True)
    stats = admin.get('/admin/stats?start=2026-09-14&end=2026-09-16')
    assert stats.status_code == 200
    assert '기간 내 예약' in stats.get_data(as_text=True)
    assert '방문 학생 수' in stats.get_data(as_text=True)
    assert '기간 설정' in stats.get_data(as_text=True)
    assert 'conic-gradient' in stats.get_data(as_text=True)
    all_stats = admin.get('/admin/stats?scope=all')
    assert all_stats.status_code == 200
    assert '전체 기간' in all_stats.get_data(as_text=True)
    assert '방문 횟수 <strong>2</strong>' in all_stats.get_data(as_text=True)
    assert '전체 방문 학생 수</span><strong>1</strong>' in all_stats.get_data(as_text=True)
    exported = admin.get('/admin/stats/export.xlsx?start=2026-09-14&end=2026-09-16')
    assert exported.status_code == 200
    workbook = load_workbook(BytesIO(exported.data), read_only=True)
    assert {'요약', '일별 예약', '학과별 회원', '학생', '예약'} <= set(workbook.sheetnames)
    assert workbook['예약'].max_row == 2
    all_export = admin.get('/admin/stats/export.xlsx?scope=all')
    assert all_export.status_code == 200
    assert load_workbook(BytesIO(all_export.data), read_only=True)['예약'].max_row == 3
    assert workbook['학생']['B2'].value.startswith("'=HYPERLINK")
    assert not any('password' in str(cell.value).lower() for row in workbook['학생'].rows for cell in row)


def test_usage_guide_is_separate_page(app):
    visitor = app.test_client()
    home = visitor.get('/').get_data(as_text=True)
    assert 'href="/usage"' in home
    assert '<details class="usage-details">' not in home
    guide = visitor.get('/usage')
    assert guide.status_code == 200
    assert '상세 이용 안내' in guide.get_data(as_text=True)
